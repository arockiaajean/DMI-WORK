import os
import random
from openpyxl import load_workbook
from tqdm import tqdm

ROOT_FOLDER = r"interal sheel dec 2025"
SKIP_KEYWORD = "consolidat"  # covers consolidate / consolidated / misspellings

# Collect all Excel files first (for file-level progress bar)
excel_files = []
for root, dirs, files in os.walk(ROOT_FOLDER):
    for file in files:
        if file.lower().endswith(".xlsx"):
            excel_files.append(os.path.join(root, file))

for file_path in tqdm(excel_files, desc="Processing Excel files"):
    wb = load_workbook(file_path)

    for sheet_name in tqdm(wb.sheetnames, desc="Sheets", leave=False):
        if SKIP_KEYWORD in sheet_name.lower():
            continue

        ws = wb[sheet_name]

        for row in tqdm(
            range(2, ws.max_row + 1),
            desc=f"Rows ({sheet_name})",
            leave=False
        ):
            student_no = ws.cell(row=row, column=2).value  # Column B

            if student_no is None:
                continue

            student_no = str(student_no)

            if (
                len(student_no) == 11 and
                student_no.isdigit() and
                student_no[8] == "4"
            ):
                ws.cell(row=row, column=4).value = random.randint(7, 10)  # D
                ws.cell(row=row, column=5).value = random.randint(7, 10)  # E
                ws.cell(row=row, column=6).value = random.randint(7, 10)  # F

    wb.save(file_path)

print("🎉 All files processed with progress bars.")

