import pandas as pd
from openpyxl import load_workbook
# with comment
# Load workbook
wb = load_workbook("TEMP ACC.xlsx")
ws = wb.active

data = []

for row in ws.iter_rows(min_row=2):  # assuming first row is header
    student_no = row[0].value
    student_name = row[1].value
    
    # iterate over remaining columns
    for cell in row[2:]:
        value = cell.value if cell.value is not None else ""
        comment = cell.comment.text if cell.comment else ""
        
        # Append both value and comment
        if value or comment:  # skip empty cells
            data.append([student_no, student_name, value, comment])

# Create DataFrame
df_long = pd.DataFrame(data, columns=["student_no", "student_name", "value", "comment"])

# Save to Excel
df_long.to_excel("cleaned_results_with_comments.xlsx", index=False)

print("✅ Finished reshaping with comments. Saved as cleaned_results_with_comments.xlsx")

