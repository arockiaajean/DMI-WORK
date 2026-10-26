import os  # REMOVE KITWE STUDENTS
import re
from openpyxl import load_workbook

ROOT_FOLDER = r"interal sheel dec 2025 1"

# Regex for exactly 11 digits
pattern = re.compile(r"^\d{11}$")

for root, dirs, files in os.walk(ROOT_FOLDER):
    for file in files:
        if file.lower().endswith(".xlsx"):
            file_path = os.path.join(root, file)

            try:
                wb = load_workbook(file_path)
                
                for sheet in wb.worksheets:

                    # Start from row 1
                    for row in range(1, sheet.max_row + 1):
                        cell_value = sheet[f"B{row}"].value

                        if cell_value is None:
                            continue

                        value_str = str(cell_value).strip()

                        if pattern.match(value_str):
                            # Hide rows where 9th character is NOT 4
                            if value_str[8] != "4":
                                sheet.row_dimensions[row].hidden = True

                wb.save(file_path)
                print(f"Processed: {file_path}")

            except Exception as e:
                print(f"Error in file {file_path}: {e}")
