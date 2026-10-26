import os
from openpyxl import load_workbook
from tqdm import tqdm

# ============================================================
# CONFIGURATION
# ============================================================

root_folder = r"C:\Users\USER\python\temp\2025\New folder"

# ============================================================
# FIND ALL EXCEL FILES
# ============================================================

excel_files = []

for foldername, subfolders, filenames in os.walk(root_folder):
    for file in filenames:
        if file.lower().endswith(".xlsx"):
            excel_files.append(os.path.join(foldername, file))

print(f"Total Excel files found: {len(excel_files)}")

# ============================================================
# PROCESS FILES
# ============================================================

successful = 0
skipped = 0

for file_path in tqdm(excel_files, desc="Processing Excel files", unit="file"):

    try:
        # ----------------------------------------------------
        # Try to open workbook with formulas
        # ----------------------------------------------------
        wb_formula = load_workbook(
            file_path,
            data_only=False
        )

        # ----------------------------------------------------
        # Try to open workbook with calculated values
        # ----------------------------------------------------
        wb_values = load_workbook(
            file_path,
            data_only=True
        )

        # ----------------------------------------------------
        # Process every worksheet
        # ----------------------------------------------------
        for sheet_name in wb_formula.sheetnames:

            ws_formula = wb_formula[sheet_name]
            ws_values = wb_values[sheet_name]

            for row in ws_formula.iter_rows():

                for cell in row:

                    # Check if cell contains a formula
                    if (
                        cell.value
                        and isinstance(cell.value, str)
                        and cell.value.startswith("=")
                    ):

                        # Get the calculated value
                        calculated_value = ws_values[cell.coordinate].value

                        # Replace formula with value
                        cell.value = calculated_value

        # ----------------------------------------------------
        # Save the workbook
        # ----------------------------------------------------
        wb_formula.save(file_path)

        successful += 1

    except Exception as e:

        # ----------------------------------------------------
        # Corrupted/unreadable file
        # Skip it and continue
        # ----------------------------------------------------
        skipped += 1

        print(f"\nSKIPPED: {file_path}")
        print(f"Reason: {e}")

        continue


# ============================================================
# FINAL SUMMARY
# ============================================================

print("\n" + "=" * 60)
print("PROCESSING COMPLETED")
print("=" * 60)

print(f"Total files found : {len(excel_files)}")
print(f"Successfully done : {successful}")
print(f"Skipped/corrupted : {skipped}")
