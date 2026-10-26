import os
import pandas as pd
import openpyxl
from copy import copy

# ================= CONFIGURATION =================
# Update these paths
MAIN_FOLDER_PATH = r"interal sheel dec 2025"
IMPACT_FILE_PATH = r"IMPACT INTERNAL FILE.xlsx"
# =================================================

def find_file_path(root_folder, target_filename):
    target_str = str(target_filename).lower().strip()
    for root, dirs, files in os.walk(root_folder):
        for file in files:
            if file.endswith((".xlsx", ".xls")):
                name_no_ext = os.path.splitext(file)[0].lower().strip()
                if name_no_ext == target_str:
                    return os.path.join(root, file)
    return None

def process_files():
    print("--- Starting Process (Appending New Row to Bottom) ---")
    
    try:
        # Load Master List (A=Student No, B=Name, C=File Name)
        df_master = pd.read_excel(IMPACT_FILE_PATH, header=0)
    except Exception as e:
        print(f"Error reading Impact Internal File: {e}")
        return

    for index, row in df_master.iterrows():
        student_no = row.iloc[0]
        student_name = row.iloc[1]
        target_file_name = row.iloc[2]
        
        if pd.isna(target_file_name):
            continue

        print(f"Processing File ID: {target_file_name}...")
        found_path = find_file_path(MAIN_FOLDER_PATH, target_file_name)

        if found_path:
            try:
                wb = openpyxl.load_workbook(found_path)
                
                # Loop through ALL sheets
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    
                    # 1. Identify the last row with data
                    last_row = ws.max_row
                    new_row = last_row + 1
                    
                    # 2. Insert Student Data into Col B (2) and C (3) of the NEW row
                    ws.cell(row=new_row, column=2).value = student_no
                    ws.cell(row=new_row, column=3).value = student_name
                    
                    # 3. "Drag" the other information
                    # We look at the row above (last_row). If it has data in other columns, copy it down.
                    if last_row > 0:
                        for col in range(1, ws.max_column + 1):
                            # Skip Column B (2) and C (3) as we just filled them
                            if col in [2, 3]:
                                continue
                            
                            # Get source cell (row above)
                            source_cell = ws.cell(row=last_row, column=col)
                            # Get target cell (new row)
                            target_cell = ws.cell(row=new_row, column=col)
                            
                            # Copy value
                            target_cell.value = source_cell.value
                            
                            # Optional: Copy number format (helps keep dates/currencies correct)
                            if source_cell.has_style:
                                target_cell.number_format = source_cell.number_format

                wb.save(found_path)
                print(f" -> Appended data to ALL sheets in: {found_path}")
                
            except Exception as e:
                print(f" -> Error editing file {found_path}: {e}")
        else:
            print(f" -> File '{target_file_name}' NOT FOUND.")

    print("--- Task Completed ---")

if __name__ == "__main__":
    process_files()
