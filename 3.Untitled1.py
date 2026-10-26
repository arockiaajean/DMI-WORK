import os
from openpyxl import load_workbook

root_folder = r"C:\Users\USER\Downloads\New folder\4-KITWE - Copy"  # <-- change this

for foldername, subfolders, filenames in os.walk(root_folder):
    for file in filenames:
        if file.endswith(".xlsx"):
            file_path = os.path.join(foldername, file)
            print(f"Processing: {file_path}")

            try:
                # Load workbook twice:
                wb_formula = load_workbook(file_path, data_only=False)
                wb_values = load_workbook(file_path, data_only=True)

                for sheet_name in wb_formula.sheetnames:
                    ws_formula = wb_formula[sheet_name]
                    ws_values = wb_values[sheet_name]

                    for row in ws_formula.iter_rows():
                        for cell in row:
                            if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                                # Replace formula with value
                                cell.value = ws_values[cell.coordinate].value

                wb_formula.save(file_path)

            except Exception as e:
                print(f"Error: {e}")
