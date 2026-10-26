from openpyxl import load_workbook, Workbook
from openpyxl.comments import Comment
from copy import copy
import os

# Folder containing Excel files
folder_path = r"C:\Users\USER\Downloads\temp"   # CHANGE THIS PATH

# Create output workbook
merged_wb = Workbook()
merged_ws = merged_wb.active
merged_ws.title = "Merged_Data"

current_row = 1
first_file = True

# Loop through all Excel files
for file_name in os.listdir(folder_path):

    # Process only .xlsx files
    if file_name.endswith(".xlsx") and file_name != "merged_output.xlsx":

        file_path = os.path.join(folder_path, file_name)

        print(f"Processing: {file_name}")

        # Load workbook
        wb = load_workbook(file_path)
        ws = wb.active

        # Skip header row after first file
        start_row = 1
        if not first_file:
            start_row = 2

        # Copy rows
        for row in ws.iter_rows(min_row=start_row):

            for cell in row:

                # Create new cell
                new_cell = merged_ws.cell(
                    row=current_row,
                    column=cell.column,
                    value=cell.value
                )

                # -------------------------
                # Copy styles
                # -------------------------
                if cell.has_style:
                    new_cell.font = copy(cell.font)
                    new_cell.fill = copy(cell.fill)
                    new_cell.border = copy(cell.border)
                    new_cell.alignment = copy(cell.alignment)
                    new_cell.number_format = copy(cell.number_format)
                    new_cell.protection = copy(cell.protection)

                # -------------------------
                # Copy comments / notes
                # -------------------------
                if cell.comment:
                    new_cell.comment = Comment(
                        text=cell.comment.text,
                        author=cell.comment.author
                    )

                # -------------------------
                # Copy hyperlinks
                # -------------------------
                if cell.hyperlink:
                    new_cell._hyperlink = copy(cell.hyperlink)

            current_row += 1

        # -------------------------
        # Copy column widths
        # -------------------------
        for col in ws.column_dimensions:
            merged_ws.column_dimensions[col].width = ws.column_dimensions[col].width

        # -------------------------
        # Copy row heights
        # -------------------------
        for row_num, row_dim in ws.row_dimensions.items():
            merged_ws.row_dimensions[row_num].height = row_dim.height

        first_file = False

# Save merged workbook
output_file = os.path.join(folder_path, "merged_output.xlsx")
merged_wb.save(output_file)

print("\nAll files merged successfully!")
print(f"Saved as: {output_file}")
