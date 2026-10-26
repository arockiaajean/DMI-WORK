from openpyxl import load_workbook

# Excel file
input_file = r"C:\Users\USER\Downloads\Sem 5 DONE\New folder\GPA\UG 2026\Business Administration-2327.xlsx"
output_file = r"temp.xlsx"

# Load workbook
wb = load_workbook(input_file)

# Process every sheet
for ws in wb.worksheets:
    # Insert a new row at the top
    ws.insert_rows(1)

    # Put the sheet name in A1
    ws["A1"] = ws.title

# Save the modified workbook
wb.save(output_file)

print("Done! All sheet names have been added to the first row.")
