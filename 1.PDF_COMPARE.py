
import fitz
import difflib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ==========================================
# 1. FILE PATHS - CHANGE THESE
# ==========================================

pdf1_path = "CMM 2.pdf"
pdf2_path = "CMM 3.pdf"

output_path = "PDF_Changes.xlsx"


# ==========================================
# 2. OPEN BOTH PDF FILES
# ==========================================

pdf1 = fitz.open(pdf1_path)
pdf2 = fitz.open(pdf2_path)

print("PDF 1 pages:", len(pdf1))
print("PDF 2 pages:", len(pdf2))

# Create Excel workbook
wb = Workbook()
ws = wb.active
ws.title = "PDF Changes"

ws.append([
    "Page Number",
    "Change Type",
    "Text in PDF 1",
    "Text in PDF 2"
])

# ==========================================
# 3. COMPARE PAGE BY PAGE
# ==========================================

total_pages = max(len(pdf1), len(pdf2))
total_changes = 0

for page_index in range(total_pages):

    # Handle pages missing from either PDF
    if page_index >= len(pdf1):
        ws.append([
            page_index + 1,
            "Page added in PDF 2",
            "",
            pdf2[page_index].get_text()
        ])
        total_changes += 1
        continue

    if page_index >= len(pdf2):
        ws.append([
            page_index + 1,
            "Page missing in PDF 2",
            pdf1[page_index].get_text(),
            ""
        ])
        total_changes += 1
        continue

    # Extract text from the current page only
    text1 = pdf1[page_index].get_text("text")
    text2 = pdf2[page_index].get_text("text")

    # Split text into lines
    lines1 = text1.splitlines()
    lines2 = text2.splitlines()

    # Compare the lines in both pages
    matcher = difflib.SequenceMatcher(
        None,
        lines1,
        lines2,
        autojunk=False
    )

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():

        # Identical text: ignore it
        if tag == "equal":
            continue

        old_text = "\n".join(lines1[i1:i2])
        new_text = "\n".join(lines2[j1:j2])

        if tag == "replace":
            change_type = "Text changed"

        elif tag == "delete":
            change_type = "Text removed from PDF 2"

        elif tag == "insert":
            change_type = "Text added in PDF 2"

        ws.append([
            page_index + 1,
            change_type,
            old_text,
            new_text
        ])

        total_changes += 1

    # Progress indicator
    if (page_index + 1) % 100 == 0:
        print("Compared", page_index + 1, "pages")

# ==========================================
# 4. FORMAT THE EXCEL FILE
# ==========================================

# Style the header
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(
        "solid",
        fgColor="1F4E78"
    )
    cell.alignment = Alignment(wrap_text=True)

# Set useful column widths
column_widths = {
    "A": 14,
    "B": 28,
    "C": 65,
    "D": 65
}

for column, width in column_widths.items():
    ws.column_dimensions[column].width = width

# Wrap long text and align it to the top
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(
            wrap_text=True,
            vertical="top"
        )

# Enable filters and keep the headings visible
ws.auto_filter.ref = ws.dimensions
ws.freeze_panes = "A2"

# Save the Excel file
wb.save(output_path)

pdf1.close()
pdf2.close()

print("\nComparison completed!")
print("Total differences found:", total_changes)
print("Excel file saved to:", output_path)

