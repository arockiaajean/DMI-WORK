import pandas as pd
from docx import Document
from docx.shared import Pt
import os
import re
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl

# --- 1. FILE CONFIGURATION ---
TEMPLATE_FILE = r"TEMP ATTENDENCE CAT.docx"
STUDENT_EXCEL = r"STUDENT_EXCEL KITWE AUGUST 2026.xlsx"
MODULE_EXCEL = r"MODULE DETAILS (2).xlsx"

OUTPUT_FOLDER = r"Generated_Attendance_Sheets KITWE AUG 2026-END 1"
if not os.path.exists(OUTPUT_FOLDER):
    os.makedirs(OUTPUT_FOLDER)

# --- HELPER FUNCTION TO CLEAN DATA ---
def clean_value(value):
    val_str = str(value).strip()
    if not val_str:
        return ""
    if val_str.endswith('.0'):
        return val_str[:-2]
    return val_str

# --- REMOVE EMPTY PAGES ---
def remove_empty_pages(doc):
    for p in doc.paragraphs[::-1]:
        if p.text.strip() == "":
            p._element.getparent().remove(p._element)
        else:
            break

# --- 2. MAIN FUNCTION ---
def create_attendance_sheets():
    print("Loading Excel files...")
    try:
        df_students = pd.read_excel(STUDENT_EXCEL)
        df_modules = pd.read_excel(MODULE_EXCEL)

        df_students['SEMESTER'] = df_students['SEMESTER'].apply(clean_value)
        df_students['PROGRAM CODE'] = df_students['PROGRAM CODE'].apply(clean_value)

        df_modules['SEMESTER'] = df_modules['SEMESTER'].apply(clean_value)
        df_modules['PROGRAM CODE'] = df_modules['PROGRAM CODE'].apply(clean_value)

    except Exception as e:
        print(f"Error reading Excel files: {e}")
        return

    count = 0

    for index, row in df_modules.iterrows():
        try:
            mod_code = str(row['MODULE CODE'])
            mod_name = str(row['MODULE NAME'])
            prog_code = clean_value(row['PROGRAM CODE'])
            prog_name = str(row['PROGRAM NAME'])
            semester = clean_value(row['SEMESTER'])

            class_list = df_students[
                (df_students['PROGRAM CODE'] == prog_code) &
                (df_students['SEMESTER'] == semester)
            ]

            if class_list.empty:
                print(f"Skipping {mod_code}: No students found.")
                continue

            # Load Word Template
            doc = Document(TEMPLATE_FILE)

            # --- HEADER REPLACEMENT ---
            replacements = {
                r"Program\s*:\s*XX": f"Program : {prog_name}",
                r"SEM\s*:\s*XX": f"SEM : {semester}",
                r"C & N:\s*XX": f"C & N: {mod_code} - {mod_name}",
                r"Session &Time:\s*XX": "Session &Time:         HRs"
            }

            for paragraph in doc.paragraphs:
                text = paragraph.text
                for pattern, replacement in replacements.items():
                    text = re.sub(pattern, replacement, text, flags=re.IGNORECASE)
                paragraph.text = text

            # --- 4. TABLE PROCESSING ---
            table = doc.tables[0]
            total_data_rows = len(table.rows) - 1

            for idx, (_, student) in enumerate(class_list.iterrows()):
                if idx < total_data_rows:
                    row_cells = table.rows[idx + 1].cells
                else:
                    row_cells = table.add_row().cells

                row_cells[0].text = str(idx + 1)
                row_cells[1].text = str(student['STUDENT NO'])
                row_cells[2].text = str(student['STUDENT NAME'])
                row_cells[3].text = ""
                row_cells[4].text = ""

            # Clear blank rows
            for idx in range(len(class_list), total_data_rows):
                for cell in table.rows[idx + 1].cells:
                    cell.text = ""

            # --- 5. FIXED FONT SIZE AFTER TABLE ---
            body = doc._body._element
            elements = list(body)

            found_table = False

            for el in elements:
                if isinstance(el, CT_Tbl):
                    found_table = True
                    continue

                if found_table and isinstance(el, CT_P):
                    for p in doc.paragraphs:
                        if p._element == el:
                            for run in p.runs:
                                run.font.size = Pt(10)
                            break

            # --- 6. UPDATE TOTAL ---
            total_count = len(class_list)
            for paragraph in doc.paragraphs:
                if "TOTAL: …………….." in paragraph.text:
                    paragraph.text = paragraph.text.replace("TOTAL: ……………..", f"TOTAL: {total_count}           ")

            # --- 7. SAVE ---
            filename = f"{prog_name}_Sem{semester}_{mod_code}.docx".replace(" ", "")
            doc.save(os.path.join(OUTPUT_FOLDER, filename))

            print(f"Created: {filename}")
            count += 1

        except Exception as e:
            print(f"Error processing module {mod_code}: {e}")

    print(f"Finished! {count} Word files generated.")

if __name__ == "__main__":
    create_attendance_sheets()

