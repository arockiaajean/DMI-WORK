
import re
from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Ask the user to enter the folder path
folder_path = input("Enter the folder path: ").strip()
folder = Path(folder_path)

# Check whether the folder exists
if not folder.is_dir():
    print("Error: The folder does not exist.")
    raise SystemExit

# Find numbers containing exactly 11 digits
pattern = re.compile(r"(?<![0-9])[0-9]{11}(?![0-9])")

# Store all extracted results here
results = []

# Search through the folder and its subfolders
for file_path in folder.rglob("*"):

    # Process PDF files
    if file_path.suffix.lower() == ".pdf":
        try:
            with fitz.open(file_path) as pdf:
                for page_number, page in enumerate(pdf, start=1):
                    text = page.get_text()

                    for number in pattern.findall(text):
                        results.append([
                            number,
                            file_path.name,
                            f"Page {page_number}"
                        ])

        except Exception as error:
            print(f"Could not read {file_path.name}: {error}")

    # Process DOCX files
    elif file_path.suffix.lower() == ".docx":
        try:
            doc = Document(file_path)

            for paragraph_number, paragraph in enumerate(
                doc.paragraphs, start=1
            ):
                for number in pattern.findall(paragraph.text):
                    results.append([
                        number,
                        file_path.name,
                        f"Paragraph {paragraph_number}"
                    ])

            # Also search tables inside Word documents
            for table_number, table in enumerate(doc.tables, start=1):
                for row_number, row in enumerate(table.rows, start=1):
                    for cell_number, cell in enumerate(row.cells, start=1):
                        for number in pattern.findall(cell.text):
                            results.append([
                                number,
                                file_path.name,
                                f"Table {table_number}, "
                                f"Row {row_number}, Cell {cell_number}"
                            ])

        except Exception as error:
            print(f"Could not read {file_path.name}: {error}")

# Create an Excel workbook
workbook = Workbook()
sheet = workbook.active
sheet.title = "Extracted Numbers"

# Add column headings
sheet.append(["11-Digit Number", "Source File", "Location"])

# Make headings bold
for cell in sheet[1]:
    cell.font = Font(bold=True)

# Write all extracted results
for result in results:
    sheet.append(result)

# Keep numbers as text to preserve leading zeros
for row in range(2, sheet.max_row + 1):
    sheet.cell(row=row, column=1).number_format = "@"

# Adjust column widths
for column in sheet.columns:
    max_length = max(len(str(cell.value or "")) for cell in column)
    sheet.column_dimensions[
        get_column_letter(column[0].column)
    ].width = min(max_length + 2, 60)

# Save Excel file inside the selected folder
output_file = folder / "extracted_11_digit_numbers.xlsx"
workbook.save(output_file)

print("\nExtraction completed!")
print("Total numbers found:", len(results))
print("Excel file saved at:", output_file)

