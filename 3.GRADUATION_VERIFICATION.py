#VERIFY
import fitz
from openpyxl import load_workbook, Workbook
from tqdm import tqdm

excel_path = "GV.xlsx"
pdf_path = "CMM 2.pdf"
output_path = "verification_results.xlsx"

# Read the input Excel file
source_wb = load_workbook(
    excel_path, read_only=True, data_only=True
)
source_ws = source_wb.active

# Create output Excel file
output_wb = Workbook()
output_ws = output_wb.active
output_ws.title = "Verification Results"

output_ws.append([
    "Student Number",
    "Student Name",
    "Number Status",
    "Name Status",
    "PDF Page"
])

# Open PDF once
with fitz.open(pdf_path) as pdf:
    total_rows = max(source_ws.max_row - 1, 0)

    for page_number, row in enumerate(
        tqdm(
            source_ws.iter_rows(min_row=2, values_only=True),
            total=total_rows,
            desc="Verifying certificates"
        ),
        start=1
    ):
        student_no = str(row[0] or "").strip()
        student_name = str(row[1] or "").strip()

        if not student_no:
            continue

        if page_number > len(pdf):
            number_status = "NO PDF PAGE"
            name_status = "NOT CHECKED"
        else:
            # Extract text from only the corresponding page
            page_text = pdf[page_number - 1].get_text("text")

            number_status = (
                "FOUND" if student_no in page_text
                else "NOT FOUND"
            )

            if number_status == "FOUND":
                name_status = (
                    "FOUND"
                    if student_name
                    and student_name.casefold() in page_text.casefold()
                    else "NOT FOUND"
                )
            else:
                name_status = "NOT CHECKED"

        output_ws.append([
            student_no,
            student_name,
            number_status,
            name_status,
            page_number
        ])

# Save results
output_wb.save(output_path)
source_wb.close()
output_wb.close()

print(f"\nVerification completed!")
print(f"Output saved to: {output_path}")

