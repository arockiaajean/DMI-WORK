
import csv
from openpyxl import load_workbook
from tqdm import tqdm

# File paths
student_file = r"temp.xlsx"
grade_file = r"temp 2.xlsx"
csv_file = r"all result cleaned 3.csv"

output_file = r"file for Dr.venice.csv"

# Read student numbers from the first Excel file
print("Reading student numbers...")

workbook = load_workbook(
    student_file, read_only=True, data_only=True
)
sheet = workbook.active

student_numbers = {
    str(row[0]).strip()
    for row in sheet.iter_rows(min_col=1, max_col=1, values_only=True)
    if row[0] is not None
}

workbook.close()

# Read approved grades from the second Excel file
print("Reading approved grades...")

workbook = load_workbook(
    grade_file, read_only=True, data_only=True
)
sheet = workbook.active

approved_grades = {
    str(row[0]).strip()
    for row in sheet.iter_rows(min_col=1, max_col=1, values_only=True)
    if row[0] is not None
}

workbook.close()

print("Student numbers loaded:", len(student_numbers))
print("Approved grades loaded:", len(approved_grades))

# Count CSV rows for progress tracking
with open(csv_file, "r", encoding="utf-8-sig", newline="") as file:
    total_rows = sum(1 for _ in file)

matched = 0

# Check the CSV file row by row
with open(csv_file, "r", encoding="utf-8-sig", newline="") as infile, \
     open(output_file, "w", encoding="utf-8-sig", newline="") as outfile:

    reader = csv.reader(infile)
    writer = csv.writer(outfile)

    for row in tqdm(reader, total=total_rows, desc="Checking results"):

        # Skip rows that do not have at least 7 columns
        if len(row) < 7:
            continue

        student_no = row[0].strip()
        grade = row[6].strip()

        # Both conditions must be satisfied
        if student_no in student_numbers and grade in approved_grades:
            writer.writerow(row)
            matched += 1

print("\nChecking completed!")
print("Matching rows:", matched)
print("Output saved to:", output_file)

