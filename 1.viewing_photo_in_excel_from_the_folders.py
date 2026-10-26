import os
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.utils import get_column_letter

# ==============================
# SETTINGS
# ==============================
IMAGE_FOLDER = r"F:\NOV GRADUATION 2026 PHOTOS"      # <-- Change this
OUTPUT_FILE = "Images.xlsx"

IMAGE_WIDTH = 100
IMAGE_HEIGHT = 100

# ==============================

wb = Workbook()
ws = wb.active
ws.title = "Photos"

# Headers
ws["A1"] = "File Name"
ws["B1"] = "Photo"

# Set column widths
ws.column_dimensions["A"].width = 35
ws.column_dimensions["B"].width = 18

supported = (".jpg", ".jpeg", ".png", ".bmp", ".gif")

row = 2

for file in sorted(os.listdir(IMAGE_FOLDER)):

    if file.lower().endswith(supported):

        full_path = os.path.join(IMAGE_FOLDER, file)

        # Write filename without extension
        ws.cell(row=row, column=1).value = os.path.splitext(file)[0]

        # Insert image
        img = Image(full_path)
        img.width = IMAGE_WIDTH
        img.height = IMAGE_HEIGHT

        ws.add_image(img, f"B{row}")

        # Adjust row height
        ws.row_dimensions[row].height = 80

        row += 1

wb.save(OUTPUT_FILE)

print("Done!")
print("Saved as:", OUTPUT_FILE)
