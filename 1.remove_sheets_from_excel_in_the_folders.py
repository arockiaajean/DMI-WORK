import os
from openpyxl import load_workbook
from difflib import SequenceMatcher

# Target sheet names (expected)
TARGET_NAMES = ["dummy no", "analysis"]

# Function to check similarity (handles spelling mistakes)
def is_similar(name, targets, threshold=0.7):
    name = name.lower().strip()
    for target in targets:
        ratio = SequenceMatcher(None, name, target).ratio()
        if ratio >= threshold:
            return True
    return False

# Root folder path
root_folder = r"C:\Users\USER\python\temp\2025"   # <-- change this

# Traverse folders
for foldername, subfolders, filenames in os.walk(root_folder):
    for file in filenames:
        if file.endswith(".xlsx"):
            file_path = os.path.join(foldername, file)
            print(f"Processing: {file_path}")
            
            try:
                wb = load_workbook(file_path)
                sheets_to_delete = []

                # Identify sheets to delete
                for sheet in wb.sheetnames:
                    if is_similar(sheet, TARGET_NAMES):
                        sheets_to_delete.append(sheet)

                # Delete sheets
                for sheet in sheets_to_delete:
                    print(f"Deleting sheet: {sheet}")
                    del wb[sheet]

                # Save only if changes were made
                if sheets_to_delete:
                    wb.save(file_path)

            except Exception as e:
                print(f"Error processing {file_path}: {e}")
print("done")
