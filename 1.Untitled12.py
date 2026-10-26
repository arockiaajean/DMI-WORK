

import os
import csv
from pathlib import Path
from openpyxl import load_workbook
from tqdm import tqdm


# ============================================================
# 1. CONFIGURATION
# ============================================================

# Main folder containing Excel files and subfolders
ROOT_FOLDER = Path(
    r"C:\Users\USER\Downloads\GPA\GPA Updated AUGUST 2026"
)

# Output CSV file
OUTPUT_CSV = ROOT_FOLDER / "consolidated_data.csv"

# Error log file
ERROR_LOG = ROOT_FOLDER / "errors.txt"


# Excel file extensions to process
EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}


# ============================================================
# 2. FIND ALL EXCEL FILES
# ============================================================

print("Searching for Excel files...")

excel_files = []

for file_path in ROOT_FOLDER.rglob("*"):

    # Skip temporary Excel files
    if file_path.name.startswith("~$"):
        continue

    # Skip the output CSV and other non-Excel files
    if file_path.is_file() and file_path.suffix.lower() in EXCEL_EXTENSIONS:

        excel_files.append(file_path)


# Sort files for consistent processing order
excel_files.sort()


print(f"Total Excel files found: {len(excel_files)}")


if not excel_files:
    print("No Excel files found.")
    raise SystemExit


# ============================================================
# 3. FIRST PASS: FIND MAXIMUM NUMBER OF COLUMNS
# ============================================================

# We need a consistent CSV header.
# First, find the maximum number of columns across all sheets.

print("\nFirst pass: Checking sheet dimensions...")

max_columns = 0
valid_files = []
errors = []

for file_path in tqdm(excel_files, desc="Checking files", unit="file"):

    try:

        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True
        )

        file_max_columns = 0

        for worksheet in workbook.worksheets:

            # Get the sheet's maximum column count
            sheet_columns = worksheet.max_column or 0

            if sheet_columns > file_max_columns:
                file_max_columns = sheet_columns

        if file_max_columns > max_columns:
            max_columns = file_max_columns

        valid_files.append(file_path)

        workbook.close()

    except Exception as error:

        error_message = (
            f"FIRST PASS ERROR | {file_path} | {error}"
        )

        errors.append(error_message)

        print(f"\nError: {file_path}")


print(f"Maximum columns found: {max_columns}")


if max_columns == 0:
    print("No columns found in the Excel files.")
    raise SystemExit


# ============================================================
# 4. CREATE CSV HEADER
# ============================================================

# Metadata columns identify the original file and sheet.
# Cell columns store the actual Excel data.

csv_header = [
    "Source_File",
    "Source_Path",
    "Sheet_Name",
    "Excel_Row_Number"
]

# Add Cell_1, Cell_2, Cell_3, etc.
for column_number in range(1, max_columns + 1):

    csv_header.append(f"Cell_{column_number}")


# ============================================================
# 5. SECOND PASS: READ ALL FILES AND WRITE CSV
# ============================================================

print("\nSecond pass: Consolidating data into CSV...")

processed_files = 0
processed_sheets = 0
processed_rows = 0
failed_files = 0


# newline="" is important for CSV files
# utf-8-sig helps Excel recognize UTF-8 text.

with open(
    OUTPUT_CSV,
    mode="w",
    newline="",
    encoding="utf-8-sig"
) as csv_file:

    csv_writer = csv.writer(
        csv_file,
        quoting=csv.QUOTE_MINIMAL
    )

    # Write the header only once
    csv_writer.writerow(csv_header)

    # Process one Excel file at a time
    for file_path in tqdm(
        valid_files,
        desc="Merging Excel files",
        unit="file"
    ):

        workbook = None

        try:

            # Read-only mode reduces memory usage
            workbook = load_workbook(
                filename=file_path,
                read_only=True,
                data_only=True
            )

            relative_path = str(
                file_path.relative_to(ROOT_FOLDER)
            )

            # Process every sheet in the current workbook
            for worksheet in workbook.worksheets:

                processed_sheets += 1

                # Read one row at a time
                for row_number, row in enumerate(
                    worksheet.iter_rows(values_only=True),
                    start=1
                ):

                    # Convert row to a list
                    row_values = list(row)

                    # Skip completely empty rows
                    if not any(
                        value is not None
                        for value in row_values
                    ):
                        continue

                    # Ensure all rows have the same number of columns
                    row_values.extend(
                        [None] * (max_columns - len(row_values))
                    )

                    # Create the output row
                    output_row = [
                        file_path.name,
                        relative_path,
                        worksheet.title,
                        row_number
                    ]

                    # Add Excel cell values
                    output_row.extend(row_values[:max_columns])

                    # Write immediately to CSV
                    csv_writer.writerow(output_row)

                    processed_rows += 1

            processed_files += 1

            # Release workbook resources
            workbook.close()
            workbook = None

        except Exception as error:

            failed_files += 1

            error_message = (
                f"SECOND PASS ERROR | {file_path} | {error}"
            )

            errors.append(error_message)

            print(f"\nError processing: {file_path}")

            # Close workbook if an error occurs
            if workbook is not None:

                try:
                    workbook.close()

                except Exception:
                    pass


# ============================================================
# 6. SAVE ERROR LOG
# ============================================================

with open(
    ERROR_LOG,
    mode="w",
    encoding="utf-8"
) as error_file:

    if errors:

        for error in errors:
            error_file.write(error + "\n")

    else:

        error_file.write("No errors found.\n")


# ============================================================
# 7. PRINT FINAL SUMMARY
# ============================================================

print("\n" + "=" * 60)
print("CONSOLIDATION COMPLETED")
print("=" * 60)

print(f"Total Excel files found: {len(excel_files)}")
print(f"Files processed successfully: {processed_files}")
print(f"Sheets processed: {processed_sheets}")
print(f"Data rows written: {processed_rows}")
print(f"Failed files: {failed_files}")
print(f"Maximum columns: {max_columns}")

print(f"\nOutput CSV:")
print(OUTPUT_CSV)

print(f"\nError log:")
print(ERROR_LOG)

print("\nDone!")
