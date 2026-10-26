import os
import multiprocessing
from openpyxl import load_workbook, Workbook
from zipfile import BadZipFile

input_folder = r"main campus 2026"
output_folder = r"YOUR_OUTPUT_FOLDER"

TIMEOUT_SECONDS = 120  # 2 minutes

os.makedirs(output_folder, exist_ok=True)


def process_excel(file_path, output_folder):

    try:
        wb = load_workbook(
            file_path,
            read_only=True,
            data_only=True
        )

        base_name = os.path.splitext(
            os.path.basename(file_path)
        )[0]

        for sheet_name in wb.sheetnames:

            ws = wb[sheet_name]

            # Create new workbook
            new_wb = Workbook()
            new_ws = new_wb.active
            new_ws.title = sheet_name

            # Copy rows
            for row in ws.iter_rows(values_only=True):
                new_ws.append(row)

            # Clean sheet name for filename
            safe_sheet_name = "".join(
                c if c.isalnum() or c in (" ", "_", "-")
                else "_"
                for c in sheet_name
            )

            output_file = (
                f"{base_name}_{safe_sheet_name}.xlsx"
            )

            output_path = os.path.join(
                output_folder,
                output_file
            )

            new_wb.save(output_path)

        wb.close()

        print(f"SUCCESS: {os.path.basename(file_path)}")

    except BadZipFile:
        print(f"BAD ZIP FILE: {os.path.basename(file_path)}")

    except Exception as e:
        print(f"ERROR: {os.path.basename(file_path)}")
        print(e)


if __name__ == "__main__":

    files = [
        f for f in os.listdir(input_folder)
        if f.endswith(".xlsx")
        and not f.startswith("~$")
    ]

    total = len(files)

    for index, file in enumerate(files, start=1):

        file_path = os.path.join(input_folder, file)

        print(f"\n[{index}/{total}] Processing: {file}")

        process = multiprocessing.Process(
            target=process_excel,
            args=(file_path, output_folder)
        )

        process.start()

        process.join(TIMEOUT_SECONDS)

        if process.is_alive():

            print(
                f"SKIPPED (Timeout > 2 mins): {file}"
            )

            process.terminate()
            process.join()
