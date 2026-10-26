#for graduation 
#	NRC	SUMIF - CSV Column 4	Starting Value - CSV Column 6	Ending Value - CSV Column 6	Matching Rows
#	127503/19/1	11400	01/03/2025	01/09/2026	19
#	170307/18/1	11400	01/03/2025	01/09/2026	19
#	118255/91/1	11400	01/03/2025	01/09/2026	19
#	299033/24/1	16800	01/06/2024	01/09/2026	28
#	330540/64/1	11400	01/03/2025	01/09/2026	19
#	220849/81/1	13300	01/03/2025	01/09/2026	19


import csv
from collections import defaultdict
from openpyxl import load_workbook, Workbook

# --------------------------------------------------
# 1. FILE LOCATIONS
# --------------------------------------------------

excel_file = "nrc list.xlsx"
csv_file = "ALL PMEC 30_10_2026_corrected.csv"
output_file = "PMEC DATA FOR GRADUATION.xlsx"

# Set True if the first row contains column headings
EXCEL_HAS_HEADER = True
CSV_HAS_HEADER = True


# --------------------------------------------------
# 2. STORE CSV DATA BY NRC
# --------------------------------------------------

results = defaultdict(lambda: {
    "total": 0.0,
    "start": None,
    "end": None,
    "count": 0
})


# --------------------------------------------------
# 3. READ THE CSV FILE
# --------------------------------------------------

with open(csv_file, "r", newline="", encoding="utf-8-sig") as file:

    reader = csv.reader(file)

    if CSV_HAS_HEADER:
        next(reader, None)

    for row in reader:

        # Skip rows with fewer than 6 columns
        if len(row) < 6:
            continue

        # Column 3: NRC
        nrc = row[2].strip()

        # Column 4: SUMIF value
        sum_value = row[3].strip()

        # Column 6: Starting and ending values
        range_value = row[5].strip()

        if not nrc:
            continue

        # Record the first matching column 6 value
        if results[nrc]["count"] == 0:
            results[nrc]["start"] = range_value

        # Update the ending value for every matching row
        results[nrc]["end"] = range_value

        results[nrc]["count"] += 1

        # SUMIF: Add column 4 values for this NRC
        try:
            number = float(sum_value.replace(",", ""))
            results[nrc]["total"] += number
        except ValueError:
            # Ignore blank or non-numeric values
            pass


# --------------------------------------------------
# 4. OPEN THE EXCEL FILE
# --------------------------------------------------

wb = load_workbook(
    excel_file,
    read_only=True,
    data_only=True
)

ws = wb.active


# --------------------------------------------------
# 5. CREATE THE OUTPUT EXCEL FILE
# --------------------------------------------------

out_wb = Workbook()
out_ws = out_wb.active
out_ws.title = "NRC Summary"

out_ws.append([
    "NRC",
    "SUMIF",
    "Starting Value",
    "Ending Value",
    "Matching Rows"
])


# --------------------------------------------------
# 6. MATCH EXCEL NRC WITH CSV NRC
# --------------------------------------------------

start_row = 2 if EXCEL_HAS_HEADER else 1

for row in ws.iter_rows(
    min_row=start_row,
    min_col=1,
    max_col=1,
    values_only=True
):

    nrc_value = row[0]

    if nrc_value is None:
        continue

    nrc = str(nrc_value).strip()

    # Find matching NRC in CSV data
    if nrc in results and results[nrc]["count"] > 0:

        data = results[nrc]

        out_ws.append([
            nrc,
            data["total"],
            data["start"],
            data["end"],
            data["count"]
        ])

    else:
        # NRC not found in CSV
        out_ws.append([
            nrc,
            0,
            "Not found",
            "Not found",
            0
        ])


# --------------------------------------------------
# 7. SAVE THE OUTPUT
# --------------------------------------------------

wb.close()

out_wb.save(output_file)

print("Processing completed successfully!")
print("Output file:", output_file)
print("Total NRCs processed:", out_ws.max_row - 1)

