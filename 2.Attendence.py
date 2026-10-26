import os
import win32com.client
from openpyxl import Workbook
# to find the documents which has more than 1 page
FOLDER_PATH = r"Generated_Attendance_Sheets KITWE AUG 2026-END"

word = win32com.client.Dispatch("Word.Application")
word.Visible = False

wb = Workbook()
ws = wb.active
ws.title = "Page Count"

ws["A1"] = "File Name"
ws["B1"] = "Page Count"

row = 2

print("Scanning folder:", FOLDER_PATH)

for root, dirs, files in os.walk(FOLDER_PATH):
    for file in files:
        if file.endswith(".docx") or file.endswith(".doc"):
            file_path = os.path.abspath(os.path.join(root, file))
            print("Processing:", file_path)

            try:
                doc = word.Documents.Open(file_path)
                pages = doc.ComputeStatistics(2)

                ws.cell(row=row, column=1, value=file)
                ws.cell(row=row, column=2, value=pages)

                doc.Close()
                row += 1

            except Exception as e:
                print(f"Error with {file}: {e}")

word.Quit()

wb.save("page_counts.xlsx")

print("Done!")
