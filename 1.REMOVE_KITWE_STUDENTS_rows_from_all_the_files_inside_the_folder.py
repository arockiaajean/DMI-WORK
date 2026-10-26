import os # REMOVE KITWE STUDENTS
import re
from openpyxl import load_workbook

ROOT_FOLDER = r"interal sheel dec 2025"

# Regex for exactly 11 digits
pattern = re.compile(r"^\d{11}$")

for root, dirs, files in os.walk(ROOT_FOLDER):
    for file in files:
        if file.lower().endswith(".xlsx"):
            file_path = os.path.join(root, file)

            try:
                wb = load_workbook(file_path)
                
                for sheet in wb.worksheets:
                    rows_to_delete = []

                    # Start from row 1 (includes header if any)
                    for row in range(1, sheet.max_row + 1):
                        cell_value = sheet[f"B{row}"].value

                        if cell_value is None:
                            continue

                        # Convert to string safely
                        value_str = str(cell_value).strip()

                        if pattern.match(value_str):
                            # Check 9th character (index 8)
                            if value_str[8] == "7":
                                rows_to_delete.append(row)
                            
                    # Delete from bottom to top (very important)
                    for row in reversed(rows_to_delete):
                        sheet.delete_rows(row, 1)

                wb.save(file_path)
                print(f"Processed: {file_path}")

            except Exception as e:
                print(f"Error in file {file_path}: {e}")

