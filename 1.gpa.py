import os
from openpyxl import load_workbook, Workbook

# === FILE PATHS ===
main_folder = r""
student_file = r".xlsx"
output_file = r"mpika gpa.xlsx"

# === LOAD STUDENT NUMBERS ===
student_wb = load_workbook(student_file)
student_ws = student_wb.active

student_numbers = set()

for row in student_ws.iter_rows(min_row=1, max_col=1, values_only=True):
    if row[0] is not None:
        student_numbers.add(str(row[0]).strip())

print("Loaded student numbers:", len(student_numbers))

# === CREATE OUTPUT FILE ===
output_wb = Workbook()
output_wb.remove(output_wb.active)  # remove default sheet

# === FUNCTION TO SEARCH IN SHEET ===
def sheet_contains_student(ws, student_numbers):
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell is not None and str(cell).strip() in student_numbers:
                return True
    return False

# === LOOP THROUGH FOLDER ===
for root, dirs, files in os.walk(main_folder):
    for file in files:
        if file.endswith(".xlsx"):
            file_path = os.path.join(root, file)
            print("Checking:", file_path)

            try:
                wb = load_workbook(file_path)

                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]

                    if sheet_contains_student(ws, student_numbers):
                        print(f"Match found in {file} - {sheet_name}")

                        # Copy sheet data
                        new_ws = output_wb.create_sheet(title=f"{file}_{sheet_name}")

                        for row in ws.iter_rows(values_only=True):
                            new_ws.append(row)

            except Exception as e:
                print("Error:", file_path, e)

# === SAVE OUTPUT ===
output_wb.save(output_file)

print("Done! Output saved to:", output_file)
