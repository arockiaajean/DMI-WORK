# ============================================
# HALL TICKET GENERATOR WITH PHOTO
# ============================================
# Install:
# pip install pillow openpyxl

from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook
import os

# ==============================
# FILES AND FOLDERS
# ==============================

excel_file = r"std_hall.xlsx"      # Excel file
photo_folder = r"C:\Users\USER\Pictures\graduant missing"           # Folder containing student photos
output_folder = r"hall_tickets"    # Output hall tickets

os.makedirs(output_folder, exist_ok=True)

# ==============================
# LOAD EXCEL
# ==============================

wb = load_workbook(excel_file)
ws = wb.active

# ==============================
# FONT SETTINGS
# ==============================

# Change font path if needed
font_path = "arial.ttf"

title_font = ImageFont.truetype(font_path, 40)
normal_font = ImageFont.truetype(font_path, 24)

# ==============================
# LOOP THROUGH STUDENTS
# ==============================

for row in ws.iter_rows(min_row=2, values_only=True):

    # Excel columns
    # A = Student No
    # B = Student Name
    # C = Program
    # D = Exam Center

    student_no = str(row[0])
    student_name = str(row[1])
    program = str(row[2])
    center = str(row[3])

    # ==============================
    # CREATE HALL TICKET IMAGE
    # ==============================

    width = 1000
    height = 600

    ticket = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(ticket)

    # Border
    draw.rectangle((10, 10, width-10, height-10), outline="black", width=4)

    # Title
    draw.text((300, 40), "HALL TICKET", fill="black", font=title_font)

    # Student Details
    draw.text((50, 150), f"Student No : {student_no}", fill="black", font=normal_font)
    draw.text((50, 220), f"Name       : {student_name}", fill="black", font=normal_font)
    draw.text((50, 290), f"Program    : {program}", fill="black", font=normal_font)
    draw.text((50, 360), f"Center     : {center}", fill="black", font=normal_font)

    # ==============================
    # LOAD PHOTO
    # ==============================

    # Example photo name:
    # 1001.jpg
    # 1002.png

    photo_found = False

    for ext in [".jpg", ".jpeg", ".png"]:

        photo_path = os.path.join(photo_folder, student_no + ext)

        if os.path.exists(photo_path):

            photo = Image.open(photo_path)

            # Resize photo
            photo = photo.resize((180, 220))

            # Paste photo
            ticket.paste(photo, (760, 150))

            photo_found = True
            break

    # If photo not found
    if not photo_found:
        draw.text((760, 250), "PHOTO\nNOT FOUND", fill="red", font=normal_font)

    # ==============================
    # SAVE
    # ==============================

    output_path = os.path.join(output_folder, f"{student_no}.png")
    ticket.save(output_path)

    print(f"Generated: {output_path}")

print("\nAll hall tickets generated successfully.")
