# to put messy nrc in column A
import openpyxl
import re
from pathlib import Path


# ============================================================
# 1. SELECT YOUR EXCEL FILE
# ============================================================

input_file = Path(r"CENTRE WISE ALLOCATION 200 nrc.xlsx")
output_file = Path(r"remaing 200 nrc.xlsx")


# ============================================================
# 2. NRC PATTERN
# ============================================================

nrc_pattern = re.compile(
    r'\b\d{6}/\d{2}/\d\b'
)


# ============================================================
# 3. OPEN EXCEL FILE
# ============================================================

wb = openpyxl.load_workbook(input_file)
ws = wb.active


# ============================================================
# 4. INSERT NEW COLUMN A
# ============================================================

ws.insert_cols(1)


# ============================================================
# 5. CHECK EVERY ROW
# ============================================================

for row in range(1, ws.max_row + 1):

    found_nrc = ""

    # Check all cells in the row, starting from column B
    for col in range(2, ws.max_column + 1):

        cell_value = ws.cell(row=row, column=col).value

        if cell_value is None:
            continue

        # Convert cell value to text
        text = str(cell_value)

        # Search for NRC
        match = nrc_pattern.search(text)

        if match:
            found_nrc = match.group()
            break

    # Put NRC in new Column A
    ws.cell(row=row, column=1).value = found_nrc


# ============================================================
# 6. SAVE NEW FILE
# ============================================================

wb.save(output_file)

print("Finished!")
print(f"Output file: {output_file}")
